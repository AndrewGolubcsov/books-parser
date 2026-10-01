"""Парсер каталога books.toscrape.com с выгрузкой в Excel.

Запуск:
    python parser.py                 # весь каталог (~1000 книг)
    python parser.py --pages 3       # только первые 3 страницы
    python parser.py --details       # плюс наличие, UPC и описание с карточки книги
"""

import argparse
import logging
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, asdict
from datetime import datetime
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

BASE_URL = "https://books.toscrape.com/"
RATING = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

log = logging.getLogger("parser")


@dataclass
class Book:
    title: str
    price: float
    rating: int
    in_stock: bool
    category: str = ""
    available: int | None = None
    upc: str = ""
    description: str = ""
    url: str = ""


def make_session() -> requests.Session:
    s = requests.Session()
    s.headers["User-Agent"] = "Mozilla/5.0 (portfolio demo parser)"
    # сайты иногда отвечают 5xx или рвут соединение — повторяем с паузой
    retry = Retry(total=3, backoff_factor=1, status_forcelist=[429, 500, 502, 503, 504])
    s.mount("https://", HTTPAdapter(max_retries=retry))
    return s


def get_soup(session: requests.Session, url: str) -> BeautifulSoup:
    r = session.get(url, timeout=15)
    r.raise_for_status()
    # сайт отдаёт utf-8, но в заголовке кодировку не указывает, и requests гадает неправильно
    r.encoding = "utf-8"
    return BeautifulSoup(r.text, "html.parser")


def parse_list_page(soup: BeautifulSoup, page_url: str) -> list[Book]:
    books = []
    for card in soup.select("article.product_pod"):
        a = card.select_one("h3 a")
        price = card.select_one(".price_color").text.strip().lstrip("Â£")
        rating_cls = card.select_one(".star-rating")["class"][1]
        books.append(Book(
            title=a["title"],
            price=float(price),
            rating=RATING.get(rating_cls, 0),
            in_stock="In stock" in card.select_one(".availability").text,
            url=urljoin(page_url, a["href"]),
        ))
    return books


def fill_details(session: requests.Session, book: Book) -> Book:
    try:
        soup = get_soup(session, book.url)
    except requests.RequestException as e:
        log.warning("не удалось открыть %s: %s", book.url, e)
        return book

    book.category = soup.select("ul.breadcrumb li")[2].text.strip()
    table = {row.th.text: row.td.text for row in soup.select("table.table tr")}
    book.upc = table.get("UPC", "")
    # «In stock (22 available)» -> 22
    digits = "".join(ch for ch in table.get("Availability", "") if ch.isdigit())
    book.available = int(digits) if digits else 0
    desc = soup.select_one("#product_description + p")
    book.description = desc.text.strip() if desc else ""
    return book


def crawl(max_pages: int | None, details: bool, workers: int) -> list[Book]:
    session = make_session()
    url = urljoin(BASE_URL, "catalogue/page-1.html")
    books: list[Book] = []
    page = 0

    while url and (max_pages is None or page < max_pages):
        page += 1
        soup = get_soup(session, url)
        books += parse_list_page(soup, url)
        log.info("страница %d: всего %d книг", page, len(books))
        nxt = soup.select_one("li.next a")
        url = urljoin(url, nxt["href"]) if nxt else None
        time.sleep(0.3)  # не долбим сайт

    if details:
        log.info("собираю карточки книг в %d потоков...", workers)
        with ThreadPoolExecutor(workers) as pool:
            books = list(pool.map(lambda b: fill_details(session, b), books))

    return books


def save_xlsx(books: list[Book], path: str, details: bool) -> None:
    columns = [
        ("title", "Название", 50),
        ("price", "Цена, £", 10),
        ("rating", "Рейтинг", 9),
        ("in_stock", "В наличии", 11),
    ]
    if details:
        columns += [
            ("category", "Категория", 18),
            ("available", "Остаток, шт", 12),
            ("upc", "UPC", 18),
            ("description", "Описание", 80),
        ]
    columns.append(("url", "Ссылка", 40))

    wb = Workbook()
    ws = wb.active
    ws.title = "Книги"
    ws.append([c[1] for c in columns])

    for b in books:
        row = asdict(b)
        ws.append(["да" if row[k] is True else "нет" if row[k] is False else row[k]
                   for k, *_ in columns])

    header_fill = PatternFill("solid", fgColor="305496")
    for i, (_, _, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # второй лист — короткая сводка, заказчики такое любят
    s = wb.create_sheet("Сводка")
    prices = [b.price for b in books]
    s.append(["Всего книг", len(books)])
    s.append(["Средняя цена, £", round(sum(prices) / len(prices), 2) if prices else 0])
    s.append(["Мин. цена, £", min(prices, default=0)])
    s.append(["Макс. цена, £", max(prices, default=0)])
    for r in range(1, 6):
        s.append([f"Книг с рейтингом {r}", sum(b.rating == r for b in books)])
    s.append(["Дата выгрузки", datetime.now().strftime("%d.%m.%Y %H:%M")])
    s.column_dimensions["A"].width = 24
    s.column_dimensions["B"].width = 18

    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description="Парсер books.toscrape.com -> Excel")
    ap.add_argument("--pages", type=int, default=None, help="сколько страниц каталога обойти")
    ap.add_argument("--details", action="store_true", help="заходить в карточку каждой книги")
    ap.add_argument("--workers", type=int, default=8, help="потоков для карточек")
    ap.add_argument("-o", "--output", default="books.xlsx")
    args = ap.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(message)s", datefmt="%H:%M:%S")

    started = time.perf_counter()
    books = crawl(args.pages, args.details, args.workers)
    save_xlsx(books, args.output, args.details)
    log.info("готово: %d книг -> %s за %.1f с", len(books), args.output, time.perf_counter() - started)


if __name__ == "__main__":
    main()
